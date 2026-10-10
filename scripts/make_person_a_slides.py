"""Generate Person A's PowerPoint slides (problem + dataset + EDA, paper-aligned).

Reference: Du, He, Zhechev - "Forecasting Bike Rental Demand" (RMSLE metric,
peak-hour buckets, temp-vs-atemp collinearity, tree models win).

Run:
    .venv/Scripts/python scripts/make_person_a_slides.py
Output:
    docs/person_a_slides.pptx
"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "person_a_slides.pptx"
OUT.parent.mkdir(exist_ok=True)

BODY_SIZE = Pt(20)


def add_title_slide(prs: Presentation, title: str, subtitle: str) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = title
    slide.placeholders[1].text = subtitle


def add_bullets_slide(prs: Presentation, title: str, bullets: list[str]) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title
    body = slide.placeholders[1].text_frame
    body.clear()
    for i, item in enumerate(bullets):
        p = body.paragraphs[0] if i == 0 else body.add_paragraph()
        p.text = item
        p.font.size = BODY_SIZE


def add_image_slide(prs: Presentation, title: str, image_name: str, caption: str) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = title
    img_path = ROOT / "reports" / image_name
    if not img_path.exists():
        print(f"warning: missing image {img_path}")
        return
    slide.shapes.add_picture(str(img_path), Inches(0.9), Inches(1.6), width=Inches(7.8))
    caption_box = slide.shapes.add_textbox(Inches(0.9), Inches(6.7), Inches(7.8), Inches(0.6))
    tf = caption_box.text_frame
    tf.text = caption
    tf.paragraphs[0].font.size = Pt(14)


def main() -> None:
    prs = Presentation()
    prs.slide_width = Inches(9.6)
    prs.slide_height = Inches(7.5)

    add_title_slide(
        prs,
        "Bike Rental Demand Forecasting",
        "Person A - Dhruv Talavat (PES1UG24AM087)\n"
        "Problem - Dataset - EDA - Preprocessing - Baseline\n"
        "Featurization aligned to Du, He & Zhechev (reference paper)",
    )

    add_bullets_slide(
        prs,
        "Problem Statement",
        [
            "Bike-share operators must match hourly supply to demand.",
            "Too few bikes at rush hour = lost rides; too many = wasted docks.",
            "Goal: predict hourly rentals (cnt) from calendar + weather features.",
            "Supervised regression on 2 years of hourly data (2011-2012).",
            "Primary metric: RMSLE (log-scale) - same as the Kaggle competition",
            "   in the reference paper; also report RMSE / MAE / R2.",
        ],
    )

    add_bullets_slide(
        prs,
        "Dataset Overview",
        [
            "Source: UCI Bike Sharing Dataset - hourly file (17,379 rows).",
            "Period: 2011-01-01 to 2012-12-31, Capital Bikeshare (Washington D.C.).",
            "Target: cnt = casual + registered rentals per hour.",
            "Calendar features: season, yr, mnth, hr, weekday, holiday, workingday.",
            "Weather: weathersit, temp, atemp, hum, windspeed (normalized).",
            "Note: casual + registered == cnt, so both are excluded (leakage).",
        ],
    )

    add_image_slide(
        prs,
        "EDA - Rentals by Hour of Day",
        "03_hourly_box.png",
        "Bimodal commute pattern: peaks at 8 AM and 5-6 PM on working days.",
    )

    add_image_slide(
        prs,
        "EDA - Weekday vs Weekend Peak Hours (paper-aligned)",
        "12_workingday_hourly_profile.png",
        "Validates the paper's peak-hour idea: weekday 7-9 AM & 5-7 PM; weekend 10 AM-6 PM.",
    )

    add_image_slide(
        prs,
        "EDA - Feels-like Temperature & Humidity",
        "06_atemp_scatter.png",
        "Demand rises with atemp then plateaus; humidity pulls demand down.",
    )

    add_image_slide(
        prs,
        "EDA - Correlation & Collinearity",
        "09_correlation_heatmap.png",
        "temp ~ atemp are ~0.99 correlated -> drop raw temp, keep atemp (paper's choice).",
    )

    add_image_slide(
        prs,
        "EDA - Seasonal Pattern",
        "04_season_avg.png",
        "Peak demand in Summer and Fall; Winter is the trough.",
    )

    add_bullets_slide(
        prs,
        "Preprocessing & Feature Engineering (paper-aligned)",
        [
            "Discretized atemp into 4 quantile buckets: cold/mild/warm/hot.",
            "Peak-hour buckets: weekday 7-9 AM, 5-7 PM; weekend 10 AM-6 PM.",
            "Kept month as a feature (paper: month beats coarse season).",
            "Dropped raw temp (collinear with atemp); kept atemp.",
            "Cyclic sin/cos encodings for hour, weekday, month.",
            "One-hot encoding with handle_unknown='ignore'; leakage-safe features.",
            "Temporal split by date: train 2011-01-01 to 2012-08-07, test after.",
        ],
    )

    add_bullets_slide(
        prs,
        "Baseline Models (for Person B to beat)",
        [
            "Linear Regression:  RMSLE 1.050 | RMSE 119.5 | R2 0.706",
            "Random Forest:      RMSLE 0.410 | RMSE  79.7 | R2 0.869",
            "Best RF params: max_depth=None, n_estimators=400, min_samples_leaf=1",
            "Paper's tree models: CTree CV-RMSLE 0.460, RF 0.503 - same pattern:",
            "   tree-based models dominate; plain linear models lag far behind.",
            "Artifacts: models/baseline_model.pkl + baseline_metrics.json +",
            "   baseline_diagnostics.json (params, features, split date).",
        ],
    )

    add_bullets_slide(
        prs,
        "Summary & Handoff",
        [
            "Cleaned, documented, leak-free dataset with paper-aligned features.",
            "Strong RF baseline (RMSLE 0.41) saved as the benchmark to beat.",
            "Handoff notes: severe-weather regime is extrapolation; RF cannot",
            "   extrapolate - gradient boosting (paper's GBM) may do better on",
            "   the recent-trend drift; hour is the most important feature.",
            "Demo input side lives in app/demo_app.py (Streamlit).",
        ],
    )

    prs.save(OUT)
    print(f"Saved {OUT}")


if __name__ == "__main__":
    main()
